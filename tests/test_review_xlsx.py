from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from facility_profiles.domain.schema import Role
from facility_profiles.pipeline.export_xlsx import (
    REVIEW_HEADERS,
    REVIEW_SHEET,
    export_workbook,
    plain_reason,
    render_value,
)
from facility_profiles.pipeline.run import Pipeline
from facility_profiles.review.xlsx_import import apply_decisions, read_decisions
from facility_profiles.storage.db import session_scope
from facility_profiles.storage.repository import Repository, unwrap
from tests.conftest import NOW, FakeTPro, sample_facility, sample_loads
from tests.test_pipeline import scripted_extractor


def _populate(settings, sessions):
    client = FakeTPro(sample_loads(), {900001: sample_facility()})
    Pipeline(settings, sessions, client=client, extractor=scripted_extractor(), now=NOW).run()  # type: ignore[arg-type]


def test_plain_reason_and_render():
    assert plain_reason("differs from the value a person entered").startswith("The notes disagree")
    assert plain_reason("conflicting values with material support").startswith("Two sources")
    assert plain_reason("confidence 0.73 between thresholds").startswith("The evidence is real")
    assert plain_reason("no supporting mention for 180 days").startswith("This value has not")
    assert plain_reason("something else") == "something else"
    assert render_value(True) == "yes" and render_value(None) is None
    assert render_value([{"days": ["mon"], "open": "07:00", "close": "15:00"}]) == "0700-1500 MON"


def test_export_then_import_round_trip(settings, sessions, tmp_path: Path):
    _populate(settings, sessions)
    out = tmp_path / "review.xlsx"
    with session_scope(sessions) as session:
        stats = export_workbook(session, out)
    assert stats.queue >= 1 and stats.fields > 0 and stats.facilities == 2
    wb = load_workbook(out)
    assert wb.sheetnames[:2] == ["Summary", REVIEW_SHEET]
    ws = wb[REVIEW_SHEET]
    headers = [c.value for c in ws[3]]
    assert headers == list(REVIEW_HEADERS)
    assert ws.data_validations.dataValidation[0].formula1 == '"accept,edit,reject"'
    field_col = REVIEW_HEADERS.index("Field") + 1
    target = next(
        r
        for r in range(4, ws.max_row + 1)
        if ws.cell(row=r, column=field_col).value == "booking_method"
    )
    first = ws[target]
    item_id = first[0].value
    assert ws.cell(row=4, column=field_col).value == "appointment_required"  # quick calls first
    why = first[REVIEW_HEADERS.index("Why it is here")].value
    assert "disagree" in why or "Two sources" in why

    # A reviewer fills the sheet in: edit one item, leave the rest blank, one bad row.
    ws.cell(row=target, column=REVIEW_HEADERS.index("Decision") + 1, value="Edit")
    ws.cell(row=target, column=REVIEW_HEADERS.index("Corrected value") + 1, value="phone")
    ws.cell(row=target, column=REVIEW_HEADERS.index("Reviewer") + 1, value="Dana")
    bad = ws.max_row + 1
    ws.cell(row=bad, column=1, value=999)
    ws.cell(row=bad, column=REVIEW_HEADERS.index("Decision") + 1, value="maybe")
    wb.save(out)

    decisions, errors = read_decisions(out)
    assert [d.item_id for d in decisions] == [item_id]
    assert decisions[0].decision == "edit" and decisions[0].corrected == "phone"
    assert errors == ["item 999: decision 'maybe' must be one of accept, edit, reject"]

    with session_scope(sessions) as session:
        repo = Repository(session)
        dry = apply_decisions(repo, decisions, default_reviewer=None, dry_run=True)
        assert dry.applied == {"edit": 1} and not dry.errors
        item = repo.get_review_item(item_id)
        assert item.status == "open"  # dry run changed nothing
        result = apply_decisions(repo, decisions, default_reviewer=None)
        assert result.total_applied == 1 and not result.errors
        item = repo.get_review_item(item_id)
        assert item.status == "edited" and item.decided_by == "Dana"
        fields = repo.fields(item.facility_key, Role(item.role))
        assert fields["booking_method"].state == "human_set"
        assert unwrap(fields["booking_method"].value) == "phone"
        again = apply_decisions(repo, decisions, default_reviewer="x")
        assert again.total_applied == 0 and "already edited" in again.errors[0]


def test_import_reports_missing_sheet_and_missing_reviewer(settings, sessions, tmp_path: Path):
    _populate(settings, sessions)
    out = tmp_path / "review.xlsx"
    with session_scope(sessions) as session:
        export_workbook(session, out)
    wb = load_workbook(out)
    ws = wb[REVIEW_SHEET]
    ws.cell(row=4, column=REVIEW_HEADERS.index("Decision") + 1, value="accept")
    wb.save(out)
    decisions, errors = read_decisions(out)
    assert len(decisions) == 1 and not errors
    with session_scope(sessions) as session:
        result = apply_decisions(Repository(session), decisions, default_reviewer=None)
    assert result.total_applied == 0 and "no Reviewer name" in result.errors[0]
    wb.remove(ws)
    wb.save(out)
    decisions, errors = read_decisions(out)
    assert decisions == [] and "not found" in errors[0]
