"""`profile set|summary|ask` and the facility filter on export-xlsx."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook
from typer.testing import CliRunner

from facility_profiles.cli import app
from facility_profiles.config import get_settings
from facility_profiles.domain.schema import FacilityIdentity, Role
from facility_profiles.storage.db import init_db, make_engine, session_factory, session_scope
from facility_profiles.storage.repository import Repository, unwrap

runner = CliRunner()


@pytest.fixture
def store(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    db = tmp_path / "fp.db"
    monkeypatch.setenv("TPRO_BASE_URL", "https://tpro.test")
    monkeypatch.setenv("TPRO_USERNAME", "u")
    monkeypatch.setenv("TPRO_PASSWORD", "p")
    monkeypatch.setenv("FP_DATABASE_URL", f"sqlite:///{db.as_posix()}")
    monkeypatch.setenv("FP_EXPORT_DIR", str(tmp_path / "exports"))
    monkeypatch.chdir(tmp_path)
    get_settings.cache_clear()
    engine = make_engine(f"sqlite:///{db.as_posix()}")
    init_db(engine)
    with session_scope(session_factory(engine)) as session:
        repo = Repository(session)
        for name, city in (
            ("Polar - Fitzgerald", "Fitzgerald"),
            ("Polar Corporation", "Worcester"),
        ):
            repo.upsert_facility(
                FacilityIdentity(
                    facility_id=None,
                    candidate_key=name.lower().replace(" ", "")[:16].ljust(16, "0"),
                    company_name=name,
                    address="1 Main St",
                    city=city,
                    state="GA",
                    postal_code="31750",
                    iana_timezone="America/New_York",
                ),
                latitude=None,
                longitude=None,
            )
    engine.dispose()
    yield db
    get_settings.cache_clear()


def test_profile_set_ask_summary_and_filtered_export(store: Path):
    # Ambiguous name is refused with the candidates listed.
    result = runner.invoke(
        app,
        ["profile", "set", "Polar", "shipper", "booking_method", "--value", "email", "--by", "t"],
    )
    assert result.exit_code != 0 and "matches 2 facilities" in result.output

    # Exact name wins; invalid values are refused; valid ones are filed as human-set.
    result = runner.invoke(
        app,
        [
            "profile",
            "set",
            "Polar - Fitzgerald",
            "shipper",
            "booking_method",
            "--value",
            "carrier-pigeon",
            "--by",
            "t",
        ],
    )
    assert result.exit_code != 0 and "not a valid value" in result.output
    result = runner.invoke(
        app,
        [
            "profile",
            "set",
            "Polar - Fitzgerald",
            "shipper",
            "appointment_required",
            "--value",
            "yes",
            "--by",
            "tester",
            "--reason",
            'stop note: "Appointments are firm"',
        ],
    )
    assert result.exit_code == 0, result.output
    assert "appointment_required = true" in result.output

    # A question goes on the queue; a later set closes it.
    result = runner.invoke(
        app,
        [
            "profile",
            "ask",
            "Polar - Fitzgerald",
            "shipper",
            "contact_email",
            "--reason",
            "address needed",
            "--proposed",
            "x@y.com",
        ],
    )
    assert result.exit_code == 0 and "contact_email: address needed" in result.output
    result = runner.invoke(
        app,
        [
            "profile",
            "set",
            "Polar - Fitzgerald",
            "shipper",
            "contact_email",
            "--value",
            "Appts@Polar.com",
            "--by",
            "megan",
        ],
    )
    assert result.exit_code == 0 and "closed 1 open review item" in result.output

    result = runner.invoke(
        app,
        [
            "profile",
            "summary",
            "Polar - Fitzgerald",
            "shipper",
            "--text",
            "Firm live appointments.",
            "--by",
            "tester",
        ],
    )
    assert result.exit_code == 0 and "summary updated" in result.output

    engine = make_engine(f"sqlite:///{store.as_posix()}")
    with session_scope(session_factory(engine)) as session:
        repo = Repository(session)
        fac = repo.find_facility(name="Polar - Fitzgerald")[0]
        fields = repo.fields(fac.key, Role.SHIPPER)
        assert unwrap(fields["appointment_required"].value) is True
        assert fields["appointment_required"].state == "human_set"
        assert unwrap(fields["contact_email"].value) == "appts@polar.com"
        assert repo.open_review_item(fac.key, Role.SHIPPER, "contact_email") is None
        assert repo.profile(fac.key, Role.SHIPPER).scheduling_summary == "Firm live appointments."
        actions = [e.action for e in repo.audit_for(fac.key)]
        assert actions.count("human_set") == 3
        other = repo.find_facility(name="Polar Corporation")[0]
        repo.ask_review(other.key, Role.SHIPPER, "booking_method", reason="how?", proposed="email")
    engine.dispose()

    # Filtered export contains only the named facility.
    out = store.parent / "vendors.xlsx"
    result = runner.invoke(
        app, ["export-xlsx", "--out", str(out), "--facility", "Polar - Fitzgerald"]
    )
    assert result.exit_code == 0, result.output
    assert "1 facilities" in result.output and "0 queue items" in result.output
    wb = load_workbook(out, read_only=True)
    names = {row[0] for row in wb["Facilities"].iter_rows(min_row=4, values_only=True)}
    assert names == {"Polar - Fitzgerald"}
    field_names = {row[3] for row in wb["Profile Fields"].iter_rows(min_row=4, values_only=True)}
    assert {"appointment_required", "contact_email"} <= field_names
