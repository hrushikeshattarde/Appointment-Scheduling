"""Read reviewer decisions back from the exported workbook and apply them (FR-11)."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from facility_profiles.pipeline.export_xlsx import DECISIONS, REVIEW_HEADERS, REVIEW_SHEET
from facility_profiles.review.queue import ReviewError, ReviewService
from facility_profiles.storage.repository import Repository


@dataclass(frozen=True)
class SheetDecision:
    """One filled-in row of the Review Queue sheet."""

    row: int
    item_id: int
    decision: str
    corrected: str | None
    reviewer: str | None
    notes: str | None


@dataclass
class ImportResult:
    """What the import did."""

    applied: dict[str, int] = field(default_factory=dict)
    skipped_blank: int = 0
    errors: list[str] = field(default_factory=list)

    @property
    def total_applied(self) -> int:
        """Decisions applied."""
        return sum(self.applied.values())


def read_decisions(path: Path) -> tuple[list[SheetDecision], list[str]]:
    """Parse the Review Queue sheet. Returns decisions and row-level errors."""
    wb = load_workbook(path, data_only=True, read_only=True)
    if REVIEW_SHEET not in wb.sheetnames:
        return [], [f"sheet '{REVIEW_SHEET}' not found in {path.name}"]
    ws = wb[REVIEW_SHEET]
    header_row = None
    columns: dict[str, int] = {}
    for row in ws.iter_rows(min_row=1, max_row=5):
        values = [str(c.value).strip() if c.value is not None else "" for c in row]
        if "Item" in values and "Decision" in values:
            header_row = row[0].row
            columns = {name: values.index(name) for name in REVIEW_HEADERS if name in values}
            break
    if header_row is None:
        return [], ["could not find the header row (needs 'Item' and 'Decision' columns)"]

    decisions: list[SheetDecision] = []
    errors: list[str] = []
    for row in ws.iter_rows(min_row=header_row + 1):
        cells: list[Any] = [c.value for c in row]
        row_number = row[0].row if row else header_row + 1

        def col(name: str, cells: list[Any] = cells) -> Any:
            index = columns.get(name)
            return cells[index] if index is not None and index < len(cells) else None

        raw_id = col("Item")
        raw_decision = col("Decision")
        if raw_id is None or raw_decision is None or not str(raw_decision).strip():
            continue  # blank decision: the reviewer skipped this row
        decision = str(raw_decision).strip().lower()
        try:
            item_id = int(str(raw_id).strip())
        except ValueError:
            errors.append(f"row with item '{raw_id}': item number is not a whole number")
            continue
        if decision not in DECISIONS:
            errors.append(
                f"item {item_id}: decision '{raw_decision}' must be one of {', '.join(DECISIONS)}"
            )
            continue
        corrected = col("Corrected value")
        corrected_text = str(corrected).strip() if corrected is not None else None
        if decision == "edit" and not corrected_text:
            errors.append(f"item {item_id}: decision is 'edit' but Corrected value is empty")
            continue
        reviewer = col("Reviewer")
        notes = col("Notes")
        decisions.append(
            SheetDecision(
                row=row_number,
                item_id=item_id,
                decision=decision,
                corrected=corrected_text or None,
                reviewer=str(reviewer).strip() if reviewer else None,
                notes=str(notes).strip() if notes else None,
            )
        )
    return decisions, errors


def apply_decisions(
    repo: Repository,
    decisions: list[SheetDecision],
    *,
    default_reviewer: str | None,
    dry_run: bool = False,
) -> ImportResult:
    """Apply sheet decisions through the review service."""
    result = ImportResult()
    service = ReviewService(repo)
    for decision in decisions:
        reviewer = decision.reviewer or default_reviewer
        if not reviewer:
            result.errors.append(f"item {decision.item_id}: no Reviewer name (or pass --by)")
            continue
        item = repo.get_review_item(decision.item_id)
        if item is None:
            result.errors.append(f"item {decision.item_id}: not found in the store")
            continue
        if item.status != "open":
            result.errors.append(f"item {decision.item_id}: already {item.status}, skipped")
            continue
        if dry_run:
            result.applied[decision.decision] = result.applied.get(decision.decision, 0) + 1
            continue
        try:
            if decision.decision == "accept":
                service.accept(decision.item_id, by=reviewer)
            elif decision.decision == "edit":
                service.edit(decision.item_id, decision.corrected or "", by=reviewer)
            else:
                service.reject(decision.item_id, by=reviewer)
        except ReviewError as exc:
            result.errors.append(f"item {decision.item_id}: {exc}")
            continue
        result.applied[decision.decision] = result.applied.get(decision.decision, 0) + 1
    return result
