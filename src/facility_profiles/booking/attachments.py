"""Text from the files attached to an email, so a confirmation sent only as a PDF is read.

A facility sometimes answers with its booking confirmation as a PDF, a Word file or a
spreadsheet, and little or nothing in the email itself. Each readable attachment's text goes
under the email's own words, marked with the file's name, so the agent matches, reads and checks
it the way it reads the email: a pickup number in the PDF backs the pickup number it reports.

Read: PDF (with the ``pdf`` extra installed, and when the PDF has text rather than a scan), plain
text, CSV, calendar invites, HTML, Word (.docx) and Excel (.xlsx). A file that should be read and
cannot be (a scanned PDF, an old .doc, a picture sent as a file) is listed as unread, so a person
is asked to open it. Logos and other pictures in signatures are left alone.
"""

from __future__ import annotations

import html
import io
import re
import zipfile
from collections.abc import Iterable
from dataclasses import dataclass
from pathlib import PurePath
from typing import Any

MAX_FILE_CHARS = 4000  # a confirmation fits; a 30-page BOL is cut
MAX_TOTAL_CHARS = 8000
MAX_FILE_BYTES = 10_000_000
MAX_UNZIPPED_BYTES = 20_000_000  # a .docx or .xlsx part larger than this is not opened
MAX_PDF_PAGES = 10
MAX_SHEET_ROWS = 200
# A picture smaller than this is a logo or a signature banner, never the confirmation.
SMALL_IMAGE_BYTES = 50_000
# A PDF with pictures and fewer words than this (links aside) is a scan or a printed picture.
PICTURE_WORDS = 25
READ = "read"

DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
_BY_EXTENSION = {
    ".pdf": "pdf",
    ".txt": "text",
    ".csv": "text",
    ".ics": "text",
    ".htm": "html",
    ".html": "html",
    ".docx": "docx",
    ".xlsx": "xlsx",
}
_BY_MIME = {
    "application/pdf": "pdf",
    "text/plain": "text",
    "text/csv": "text",
    "text/calendar": "text",
    "application/ics": "text",
    "text/html": "html",
    DOCX: "docx",
    XLSX: "xlsx",
}
# Files that carry nothing to read: mail signatures and contact cards.
_SKIP_MIMES = frozenset(
    {
        "application/pkcs7-signature",
        "application/x-pkcs7-signature",
        "application/pgp-signature",
        "text/vcard",
        "text/x-vcard",
        "text/directory",
    }
)
_SKIP_EXTENSIONS = frozenset({".p7s", ".asc", ".sig", ".vcf"})
_SPACE_RE = re.compile(r"[ \t\r\f\v]+")
_URL_RE = re.compile(r"\b(?:https?://|www\.)\S+", re.I)
_BLANK_LINES_RE = re.compile(r"\n\s*\n\s*\n+")


@dataclass(frozen=True)
class Attachment:
    """One attached file as a mail source hands it over."""

    filename: str
    mime: str
    data: bytes
    inline: bool = False  # shown in the body (a logo) rather than attached


@dataclass(frozen=True)
class Reading:
    """What came out of one file: its text, or why a person has to open it."""

    name: str
    text: str | None = None
    why: str | None = None  # set when the file should have been read and was not


def _kind(filename: str, mime: str) -> str | None:
    extension = PurePath(filename.lower()).suffix
    return _BY_EXTENSION.get(extension) or _BY_MIME.get(mime.lower().split(";")[0].strip())


def triage(filename: str, mime: str, size: int, *, inline: bool = False) -> str | None:
    """Whether to read a file before fetching it.

    :data:`READ` for a file whose text the agent takes; None for one to leave alone (a logo, a
    signature file, an empty part); otherwise why the agent cannot read it.
    """
    mime = mime.lower().split(";")[0].strip()
    extension = PurePath(filename.lower()).suffix
    if size == 0 or mime in _SKIP_MIMES or extension in _SKIP_EXTENSIONS:
        return None
    if mime.startswith("image/"):
        if inline or size < SMALL_IMAGE_BYTES:
            return None
        return "a picture; the agent cannot read pictures"
    if size > MAX_FILE_BYTES:
        return f"too large to read ({size // 1_000_000} MB)"
    if _kind(filename, mime) is not None:
        return READ
    what = f"{extension} files" if extension else "this kind of file"
    return f"the agent cannot read {what}"


def read(attachment: Attachment) -> Reading | None:
    """The file's text, or why it could not be read; None when it is not worth reading."""
    name = attachment.filename or "attachment"
    verdict = triage(
        attachment.filename, attachment.mime, len(attachment.data), inline=attachment.inline
    )
    if verdict is None:
        return None
    if verdict != READ:
        return Reading(name, why=verdict)
    kind = _kind(attachment.filename, attachment.mime)
    text, why = _extract(kind or "text", attachment.data)
    if why is not None:
        return Reading(name, why=why)
    cleaned = _tidy(text or "")
    if not cleaned:
        return Reading(name, why="the file has no text in it")
    return Reading(name, text=cleaned)


def _extract(kind: str, data: bytes) -> tuple[str | None, str | None]:
    if kind == "pdf":
        return _pdf(data)
    try:
        if kind == "docx":
            return _docx(data), None
        if kind == "xlsx":
            return _xlsx(data), None
    except (zipfile.BadZipFile, KeyError, ValueError, OSError) as exc:
        return None, f"the file could not be opened ({type(exc).__name__})"
    text = _decode(data)
    return (_html_text(text) if kind == "html" else text), None


def _decode(data: bytes) -> str:
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        return data.decode("latin-1")


def _pdf(data: bytes) -> tuple[str | None, str | None]:
    try:
        from pypdf import PdfReader  # noqa: PLC0415 - the optional pdf extra
    except ImportError:
        return None, "PDF reading is not installed here (the pdf extra)"
    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            reader.decrypt("")  # most "protected" PDFs open with an empty password
        pages = list(reader.pages[:MAX_PDF_PAGES])
        text = "\n".join(page.extract_text() or "" for page in pages)
    except Exception as exc:  # a broken or locked PDF is one a person opens
        return None, f"the PDF could not be opened ({type(exc).__name__})"
    if not text.strip():
        return None, "the PDF has no text in it (a scan or a picture)"
    words = _URL_RE.sub(" ", text)
    if len(words.split()) < PICTURE_WORDS and _has_pictures(pages):
        # A picture printed to PDF: what text there is, is the printer's header and the page's
        # address, not what the picture says.
        return None, "the PDF is a picture with next to no text (a scan or a printout)"
    return text, None


def _has_pictures(pages: list[Any]) -> bool:
    try:
        return any(len(page.images) for page in pages)
    except Exception:  # an image the reader cannot list is still an image
        return True


def _unzipped(archive: zipfile.ZipFile, name: str) -> str:
    if archive.getinfo(name).file_size > MAX_UNZIPPED_BYTES:
        msg = f"{name} is too large"
        raise ValueError(msg)
    return archive.read(name).decode("utf-8", "replace")


def _docx(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        xml = _unzipped(archive, "word/document.xml")
    paragraphs = re.split(r"</w:p>", xml)
    return "\n".join(
        html.unescape("".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", p))) for p in paragraphs
    )


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook  # noqa: PLC0415 - only when a spreadsheet comes in

    with zipfile.ZipFile(io.BytesIO(data)) as archive:  # a zip bomb is refused before openpyxl
        for info in archive.infolist():
            if info.file_size > MAX_UNZIPPED_BYTES:
                msg = f"{info.filename} is too large"
                raise ValueError(msg)
    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    lines: list[str] = []
    try:
        for sheet in book.worksheets:
            lines.append(f"[{sheet.title}]")
            for row in sheet.iter_rows(max_row=MAX_SHEET_ROWS, values_only=True):
                cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if cells:
                    lines.append(" | ".join(cells))
    finally:
        book.close()
    return "\n".join(lines)


def _html_text(text: str) -> str:
    text = re.sub(r"<(?:script|style).*?</(?:script|style)>", " ", text, flags=re.S | re.I)
    text = re.sub(r"<br\s*/?>|</p>|</div>|</tr>|</li>", "\n", text, flags=re.I)
    return html.unescape(re.sub(r"<[^>]+>", " ", text))


def _tidy(text: str) -> str:
    lines = (_SPACE_RE.sub(" ", line).strip() for line in text.replace("\r\n", "\n").split("\n"))
    return _BLANK_LINES_RE.sub("\n\n", "\n".join(lines)).strip()


def mark(name: str) -> str:
    """The line that opens a file's text under the email's own words."""
    return f"--- Attached file: {name} ---"


def with_attachments(body: str, readings: Iterable[Reading]) -> tuple[str, tuple[str, ...]]:
    """The email's own words with each read file's text under them, and the files not read.

    Each unread file is given as "name: why", and named in the text too ("--- Attached file not
    read: ..."), so whoever reads the email on the board knows to open it. A long file is cut to
    :data:`MAX_FILE_CHARS` (a confirmation is on its first page); one that finds no room left of
    :data:`MAX_TOTAL_CHARS` is listed as not read.
    """
    parts = [body.rstrip()] if body.strip() else []
    unread: list[str] = []
    room = MAX_TOTAL_CHARS
    for reading in readings:
        why = reading.why or "not read"
        if reading.text is not None:
            limit = min(MAX_FILE_CHARS, room)
            if limit > 0:
                text = reading.text[:limit]
                parts.append(f"{mark(reading.name)}\n{text}")
                room -= len(text)
                continue
            why = "more text than the agent reads in one email"
        unread.append(f"{reading.name}: {why}")
        parts.append(f"--- Attached file not read: {reading.name} ({why}) ---")
    return "\n\n".join(parts), tuple(unread)
